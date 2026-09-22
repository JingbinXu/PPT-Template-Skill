#!/usr/bin/env python3
"""Render-free QA for an editable PPTX created by the template workflow."""
from __future__ import annotations

import argparse
import json
from itertools import combinations
from pathlib import Path

from pptx import Presentation


def overlap(a, b):
    left = max(a["x"], b["x"])
    top = max(a["y"], b["y"])
    right = min(a["x"] + a["w"], b["x"] + b["w"])
    bottom = min(a["y"] + a["h"], b["y"] + b["h"])
    return max(0, right - left) * max(0, bottom - top)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    source = Path(args.input)
    prs = Presentation(source)
    width, height = prs.slide_width, prs.slide_height
    report = {"input": str(source.resolve()), "slide_count": len(prs.slides), "canvas_emu": {"width": width, "height": height}, "fatal": [], "warnings": [], "slides": [], "fonts": [], "pending_content": [], "template_decoration_outside_canvas": []}
    for number, slide in enumerate(prs.slides, 1):
        text_shapes, slide_text = [], []
        for index, shape in enumerate(slide.shapes, 1):
            item = {"index": index, "name": shape.name, "x": shape.left, "y": shape.top, "w": shape.width, "h": shape.height}
            text = ""
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name and run.font.name not in report["fonts"]:
                            report["fonts"].append(run.font.name)
            elif getattr(shape, "has_table", False):
                text = "\n".join(cell.text for row in shape.table.rows for cell in row.cells if cell.text.strip())
            if text:
                item["text"] = text[:240]
                slide_text.append(text)
                text_shapes.append(item)
                if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width or shape.top + shape.height > height:
                    report["fatal"].append({"slide": number, "shape": index, "issue": "editable_content_outside_canvas"})
                if "【待补充" in text:
                    report["pending_content"].append({"slide": number, "text": text[:240]})
            elif shape.left < 0 or shape.top < 0 or shape.left + shape.width > width or shape.top + shape.height > height:
                report["template_decoration_outside_canvas"].append({"slide": number, "shape": index})
        for a, b in combinations(text_shapes, 2):
            area = overlap(a, b)
            if area > 50000000000:
                report["warnings"].append({"slide": number, "issue": "possible_text_overlap", "shapes": [a["index"], b["index"]]})
        report["slides"].append({"slide": number, "shape_count": len(slide.shapes), "text_preview": "\n".join(slide_text)[:1000]})
    if not prs.slides:
        report["fatal"].append({"issue": "no_slides"})
    destination = Path(args.out)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(destination)
    if report["fatal"]:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
