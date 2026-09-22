#!/usr/bin/env python3
"""Write a compact, source-safe visual inventory for a PPTX template."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation


def inches(value: int) -> float:
    return round(value / 914400, 3)


def font_sample(shape):
    if not getattr(shape, "has_text_frame", False):
        return None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                color = None
                try:
                    color = run.font.color.rgb
                    color = str(color) if color else None
                except Exception:
                    pass
                return {
                    "name": run.font.name,
                    "size_pt": round(run.font.size.pt, 1) if run.font.size else None,
                    "bold": bool(run.font.bold) if run.font.bold is not None else None,
                    "color": color,
                }
    return None


def shape_item(shape, index: int):
    text = shape.text.strip() if getattr(shape, "has_text_frame", False) else ""
    if len(text) > 160:
        text = text[:157] + "..."
    return {
        "index": index,
        "name": shape.name,
        "type": str(shape.shape_type),
        "x": inches(shape.left),
        "y": inches(shape.top),
        "w": inches(shape.width),
        "h": inches(shape.height),
        "text": text,
        "font": font_sample(shape),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    input_path = Path(args.input).resolve()
    if not input_path.is_file() or input_path.suffix.lower() != ".pptx":
        raise SystemExit(f"Template must be a readable .pptx: {input_path}")
    presentation = Presentation(input_path)
    slides = []
    fonts = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        shapes = [shape_item(shape, index) for index, shape in enumerate(slide.shapes, start=1)]
        for item in shapes:
            if item["font"] and item["font"] not in fonts:
                fonts.append(item["font"])
        slides.append({
            "slide_index": slide_number,
            "shape_count": len(shapes),
            "text_samples": [item["text"] for item in shapes if item["text"]][:8],
            "shapes": shapes,
        })
    result = {
        "input": str(input_path),
        "canvas": {"width_in": inches(presentation.slide_width), "height_in": inches(presentation.slide_height)},
        "slide_count": len(slides),
        "font_samples": fonts[:12],
        "slides": slides,
    }
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(out)


if __name__ == "__main__":
    main()
